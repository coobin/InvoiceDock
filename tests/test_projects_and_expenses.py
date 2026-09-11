from datetime import UTC, datetime
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from starlette.requests import Request

from app.db import Base
from app.main import (
    app,
    expenses_page,
    invoice_query,
    owned_expense,
    owned_project,
    projects_page,
)
from app.models import ExpenseItem, Invoice, Project, User
from app.security import session_auth_marker
from app.services.export_service import make_invoice_workbook


def _factory(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'proj_exp.db'}")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, class_=Session, expire_on_commit=False)


def _request(user: User, path: str = "/projects") -> Request:
    return Request(
        {
            "type": "http",
            "http_version": "1.1",
            "method": "GET",
            "scheme": "https",
            "path": path,
            "raw_path": path.encode(),
            "query_string": b"",
            "headers": [],
            "client": ("127.0.0.1", 1),
            "server": ("testserver", 443),
            "session": {
                "user_id": user.id,
                "auth_marker": session_auth_marker(user),
            },
            "app": app,
            "router": app.router,
        }
    )


def test_project_ownership_and_isolation(tmp_path):
    factory = _factory(tmp_path)
    with factory() as db:
        user_a = User(username="user_a", email="a@example.com")
        user_b = User(username="user_b", email="b@example.com")
        admin = User(username="admin_user", email="admin@example.com", role="admin")
        db.add_all([user_a, user_b, admin])
        db.flush()

        proj_a = Project(owner_id=user_a.id, name="项目A", code="PRJ-A", budget=10000.0)
        proj_b = Project(owner_id=user_b.id, name="项目B", code="PRJ-B", budget=20000.0)
        db.add_all([proj_a, proj_b])
        db.commit()

        req_a = _request(user_a)
        assert owned_project(req_a, db, user_a, proj_a.id).name == "项目A"
        with pytest.raises(HTTPException) as exc:
            owned_project(req_a, db, user_a, proj_b.id)
        assert exc.value.status_code == 404

        # 管理员有权访问任意项目
        req_admin = _request(admin)
        assert owned_project(req_admin, db, admin, proj_b.id).name == "项目B"


def test_expense_pre_import_and_reconciliation(tmp_path):
    factory = _factory(tmp_path)
    with factory() as db:
        user = User(username="employee", email="emp@example.com")
        db.add(user)
        db.flush()

        proj = Project(owner_id=user.id, name="商务拓展", code="BD-01")
        db.add(proj)
        db.flush()

        expense = ExpenseItem(
            owner_id=user.id,
            project_id=proj.id,
            claimant="张三",
            expected_amount=520.0,
            category="餐饮",
            expense_date="2026-08-20",
            expected_seller="某某海鲜酒楼",
            notes="客户招待晚餐",
            status="pending",
        )
        db.add(expense)
        db.commit()

        # 验证待开票初始状态
        req = _request(user, "/expenses")
        fetched = owned_expense(req, db, user, expense.id)
        assert fetched.status == "pending"
        assert fetched.expected_amount == 520.0
        assert fetched.reconciled_invoice_id is None

        # 真实发票到达入库
        invoice = Invoice(
            owner_id=user.id,
            original_name="dinner_inv.pdf",
            stored_name="dinner_inv.pdf",
            sha256="hash12345",
            total_amount=535.0,
            seller_name="某某海鲜酒楼有限公司",
            invoice_number="260012345678",
            invoice_date="2026-08-20",
            status="verified",
        )
        db.add(invoice)
        db.commit()

        # 执行核销
        expense.reconciled_invoice_id = invoice.id
        expense.status = "reconciled"
        expense.reconciled_at = datetime.now(UTC).replace(tzinfo=None)
        if not invoice.project_id:
            invoice.project_id = expense.project_id
        db.commit()

        assert expense.status == "reconciled"
        assert expense.reconciled_invoice_id == invoice.id
        assert invoice.project_id == proj.id


def test_invoice_query_with_project_filter(tmp_path):
    factory = _factory(tmp_path)
    with factory() as db:
        user = User(username="finance", email="fin@example.com")
        db.add(user)
        db.flush()

        p1 = Project(owner_id=user.id, name="项目一")
        p2 = Project(owner_id=user.id, name="项目二")
        db.add_all([p1, p2])
        db.flush()

        inv1 = Invoice(owner_id=user.id, original_name="1.pdf", stored_name="1.pdf", sha256="s1", project_id=p1.id)
        inv2 = Invoice(owner_id=user.id, original_name="2.pdf", stored_name="2.pdf", sha256="s2", project_id=p2.id)
        inv3 = Invoice(owner_id=user.id, original_name="3.pdf", stored_name="3.pdf", sha256="s3", project_id=None)
        db.add_all([inv1, inv2, inv3])
        db.commit()

        q1 = list(db.scalars(invoice_query(project_id=p1.id, user=user)).all())
        assert len(q1) == 1
        assert q1[0].id == inv1.id

        q_all = list(db.scalars(invoice_query(user=user)).all())
        assert len(q_all) == 3


def test_excel_export_includes_project_and_expense_sheets():
    p = Project(name="新一代平台重构", code="PRJ-2026")
    p.id = "proj-uuid-1"

    invoice = Invoice(
        original_name="test.pdf",
        stored_name="test.pdf",
        sha256="test",
        status="verified",
        verification_method="kingdee",
        invoice_number="24500000000012345678",
        seller_name="示例供应商",
        total_amount=1200.0,
        project_id=p.id,
        created_at=datetime(2026, 8, 10, 10, 0, 0),
    )

    expense = ExpenseItem(
        claimant="李四",
        expected_amount=600.0,
        category="住宿",
        expense_date="2026-08-11",
        expected_seller="希尔顿酒店",
        expected_invoice_date="2026-08-15",
        status="pending",
        project_id=p.id,
        notes="出差住宿待开票",
        created_at=datetime(2026, 8, 10, 12, 0, 0),
    )

    wb_bytes = make_invoice_workbook(
        [invoice],
        projects_map={p.id: p},
        expenses=[expense],
        invoices_map={invoice.id: invoice},
    )
    wb = load_workbook(BytesIO(wb_bytes))

    # 检查 Sheet 1: 发票台账包含项目信息
    sheet1 = wb["发票台账"]
    assert sheet1["A2"].value == "verified"
    assert sheet1["E2"].value == "24500000000012345678"
    assert sheet1["S2"].value == "新一代平台重构"
    assert sheet1["T2"].value == "PRJ-2026"

    # 检查 Sheet 2: 待开票预录包含正确列
    assert "待开票预录" in wb.sheetnames
    sheet2 = wb["待开票预录"]
    assert sheet2["A2"].value == "待开票"
    assert sheet2["B2"].value == "新一代平台重构"
    assert sheet2["C2"].value == "李四"
    assert sheet2["E2"].value == 600.0
    assert sheet2["F2"].value == "住宿"


def test_projects_and_expenses_page_rendering(tmp_path):
    factory = _factory(tmp_path)
    with factory() as db:
        user = User(username="viewer", email="viewer@example.com")
        db.add(user)
        db.flush()

        proj = Project(owner_id=user.id, name="测试页面项目", budget=50000.0)
        db.add(proj)
        db.flush()

        exp = ExpenseItem(
            owner_id=user.id,
            project_id=proj.id,
            claimant="王五",
            expected_amount=300.0,
            status="pending",
            expense_date="2026-08-15",
        )
        db.add(exp)
        db.commit()

        # 渲染项目页面
        proj_resp = projects_page(_request(user, "/projects"), db)
        assert proj_resp.status_code == 200
        assert "测试页面项目" in proj_resp.body.decode()

        # 渲染待开票预录页面
        exp_resp = expenses_page(_request(user, "/expenses"), db=db)
        assert exp_resp.status_code == 200
        assert "王五" in exp_resp.body.decode()
        assert "300.00" in exp_resp.body.decode()
